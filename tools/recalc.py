#!/usr/bin/env python3
"""Recalculate every formula in an .xlsx with LibreOffice, in place, and report
Excel error values. Used by tests/export.test.js for the Excel-vs-app parity
check. Needs LibreOffice Calc (soffice on PATH) and openpyxl.

    python3 tools/recalc.py <file.xlsx> [timeout_seconds]

Prints JSON: {"status": "success" | "errors_found", "total_errors": n,
"total_formulas": n, "error_summary": {...}} or {"error": "..."}.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
ERRORS = ["#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"]


def soffice(args, timeout, env):
    cmd = ["soffice"] + args
    if shutil.which("timeout"):
        cmd = ["timeout", str(timeout)] + cmd
    return subprocess.run(cmd, capture_output=True, text=True, env=env, timeout=timeout + 15)


def recalc(path, timeout):
    path = Path(path).absolute()
    if not path.exists():
        return {"error": f"{path} does not exist"}
    if not shutil.which("soffice"):
        return {"error": "soffice not found on PATH; LibreOffice Calc is required"}
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    with tempfile.TemporaryDirectory(prefix="recalc-lo-") as prof:
        url = Path(prof).as_uri()
        # First start creates the profile; then the macro goes into it
        r = soffice(["--headless", "--terminate_after_init", f"-env:UserInstallation={url}"], timeout, env)
        macro_dir = Path(prof) / "user" / "basic" / "Standard"
        if not macro_dir.exists():
            return {"error": "LibreOffice did not create a profile: " + (r.stderr or "").strip()}
        (macro_dir / "Module1.xba").write_text(MACRO)
        before = (path.stat().st_mtime_ns, path.stat().st_size)
        r = soffice(["--headless", "--norestore", f"-env:UserInstallation={url}",
                     "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application", str(path)], timeout, env)
        if r.returncode == 124:
            return {"error": f"LibreOffice timed out after {timeout}s"}
        if r.returncode != 0:
            return {"error": "LibreOffice failed: " + ((r.stderr or "").strip() or f"exit {r.returncode}")}
        if (path.stat().st_mtime_ns, path.stat().st_size) == before:
            return {"error": "LibreOffice exited but did not rewrite the file"}
    values = load_workbook(path, data_only=True)
    summary, total = {}, 0
    for ws in values.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str):
                    for e in ERRORS:
                        if e in v:
                            summary.setdefault(e, {"count": 0, "locations": []})
                            summary[e]["count"] += 1
                            if len(summary[e]["locations"]) < 100:
                                summary[e]["locations"].append(f"{ws.title}!{cell.coordinate}")
                            total += 1
                            break
    values.close()
    formulas = load_workbook(path, data_only=False)
    count = 0
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, ArrayFormula):
                    v = v.text
                if isinstance(v, str) and v.startswith("="):
                    count += 1
    formulas.close()
    return {"status": "success" if total == 0 else "errors_found", "total_errors": total, "total_formulas": count, "error_summary": summary}


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    print(json.dumps(recalc(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 60)))
